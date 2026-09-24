"""Blind teacher rating: export shuffled xlsx sheets, import them back to ratings.csv."""
from __future__ import annotations

import random
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .config import resolve
from .rubric import CRITERIA, RUBRIC_HU
from .scoring import latest_exams
from .util import read_jsonl

LETTERS = "ABCD"
HU_COLS = {"correctness": "Helyesség (1-5)", "clarity": "Érthetőség (1-5)",
           "distractor_quality": "Disztraktorok (1-5)", "bloom_fit": "Szintillesztés (1-5)"}
FIXED_COLS = ["Azonosító", "Dokumentum", "Nehézség", "Forrásrészlet", "Kérdés", "A", "B", "C", "D",
              "Megjelölt helyes"]
TAIL_COLS = [*(HU_COLS[c] for c in CRITERIA), "Használnám (igen/nem)", "Megjegyzés"]
DIFF_HU = {"easy": "könnyű", "medium": "közepes", "hard": "nehéz"}


def _sample(rows: list[dict], n: int, rng: random.Random) -> list[dict]:
    """Round-robin over documents so every document is represented before any repeats."""
    by_doc = defaultdict(list)
    for r in rows:
        by_doc[r["doc_id"]].append(r)
    for v in by_doc.values():
        rng.shuffle(v)
    docs = sorted(by_doc)
    rng.shuffle(docs)
    out = []
    while len(out) < n and any(by_doc.values()):
        for d in docs:
            if by_doc[d] and len(out) < n:
                out.append(by_doc[d].pop())
    return out


def _excerpt(q: dict, exams: dict) -> str:
    ex = exams.get((q["doc_id"], q["seed"]))
    if not ex:
        return ""
    idx = q.get("nearest_chunk")
    chunks = ex.get("chunks") or []
    if idx is not None and 0 <= idx < len(chunks):
        return chunks[idx][:1500]
    return (ex.get("context_text") or "")[:1500]


def export_sheets(run_dirs: list[str], out_dir: str, n_per_run: int = 30, n_raters: int = 3,
                  n_calibration: int = 10, seed: int = 7, types: tuple = ("mcq", "tf", "open")) -> Path:
    rng = random.Random(seed)
    out = resolve(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    items, calib = [], []
    for rd in run_dirs:
        rd = resolve(rd)
        qs = [q for q in read_jsonl(rd / "questions.jsonl") if q.get("type") in types]
        if not qs:
            raise ValueError(f"{rd}: no scored questions (run `score` first)")
        exams = {(e["doc_id"], e["seed"]): e for e in latest_exams(read_jsonl(rd / "exams.jsonl"))}
        picked = _sample(qs, n_per_run, rng)
        if len(picked) < n_per_run:
            print(f"[rate-export] WARNING: {rd.name} has only {len(picked)} questions")
        rest = [q for q in qs if q["uid"] not in {p["uid"] for p in picked}]
        cal = _sample(rest, max(0, n_calibration // max(1, len(run_dirs))), rng)
        for q in picked:
            items.append({**q, "_excerpt": _excerpt(q, exams)})
        for q in cal:
            calib.append({**q, "_excerpt": _excerpt(q, exams)})
    rng.shuffle(items)
    for i, q in enumerate(items, 1):
        q["item_id"] = f"Q{i:03d}"
    for i, q in enumerate(calib, 1):
        q["item_id"] = f"K{i:02d}"

    key = pd.DataFrame([{"item_id": q["item_id"], "uid": q["uid"], "arm": q["arm"], "run_id": q["run_id"],
                         "doc_id": q["doc_id"], "calibration": q["item_id"].startswith("K")}
                        for q in calib + items])
    key.to_csv(out / "rating_key.csv", index=False)

    for r in range(n_raters):
        rater = f"rater_{chr(ord('A') + r)}"
        order = items[:]
        random.Random(seed * 100 + r).shuffle(order)   # different order per rater
        _write_workbook(out / f"{rater}.xlsx", rater, calib, order)
    print(f"[rate-export] {len(items)} items (+{len(calib)} calibration) x {n_raters} raters -> {out}\n"
          f"             keep rating_key.csv away from the raters.")
    return out


def _write_workbook(path: Path, rater: str, calib: list[dict], items: list[dict]) -> None:
    wb = Workbook()
    guide = wb.active
    guide.title = "Útmutató"
    lines = [f"Értékelő: {rater}", "",
             "Minden kérdést a mellette lévő forrásrészlet alapján értékelj. A kérdéseket egy MI-rendszer és/vagy "
             "más módszer készítette; nem tudod, melyik melyiktől származik. Ne beszéld meg a pontszámokat a "
             "többi értékelővel.", "",
             "Először a 'Kalibráció' lapot töltsd ki (ezt közösen megbeszéljük), utána az 'Értékelés' lapot.", "",
             "Pontozási útmutató:"] + [f"• {RUBRIC_HU[c]}" for c in CRITERIA] + [f"• {RUBRIC_HU['would_use']}"]
    for i, line in enumerate(lines, 1):
        guide.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
    guide.column_dimensions["A"].width = 120
    guide["A1"].font = Font(bold=True)
    if calib:
        _fill_sheet(wb.create_sheet("Kalibráció"), calib)
    _fill_sheet(wb.create_sheet("Értékelés"), items)
    wb.save(path)


def _fill_sheet(ws, items: list[dict]) -> None:
    cols = FIXED_COLS + TAIL_COLS
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for q in items:
        opts = [o["text"] for o in q.get("options", [])][:4]
        correct = [LETTERS[i] for i, o in enumerate(q.get("options", [])[:4]) if o.get("correct")]
        row = [q["item_id"], q["doc_id"], DIFF_HU.get(q.get("difficulty"), q.get("difficulty")),
               q["_excerpt"], q["text"], *(opts + [""] * (4 - len(opts))),
               ", ".join(correct) if q["type"] != "open" else f"(kifejtős) {q.get('key', '')}"]
        ws.append(row + [None] * len(TAIL_COLS))
    n = len(items) + 1
    first_score = len(FIXED_COLS) + 1
    dv15 = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                          showErrorMessage=True, error="1 és 5 közötti egész szám")
    dvyn = DataValidation(type="list", formula1='"igen,nem"', allow_blank=True)
    ws.add_data_validation(dv15)
    ws.add_data_validation(dvyn)
    for j in range(len(CRITERIA)):
        col = ws.cell(row=1, column=first_score + j).column_letter
        dv15.add(f"{col}2:{col}{n}")
    yn = ws.cell(row=1, column=first_score + len(CRITERIA)).column_letter
    dvyn.add(f"{yn}2:{yn}{n}")
    widths = [9, 16, 10, 60, 45, 25, 25, 25, 25, 12] + [13] * len(CRITERIA) + [12, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"


def import_sheets(sheets: list[str], key_csv: str, out_csv: str) -> Path:
    key = pd.read_csv(resolve(key_csv))
    rows = []
    rev = {v: k for k, v in HU_COLS.items()}
    for s in sheets:
        p = resolve(s)
        wb = load_workbook(p, data_only=True)
        rater = p.stem
        for sheet in ("Kalibráció", "Értékelés"):
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            header = [c.value for c in ws[1]]
            for vals in ws.iter_rows(min_row=2, values_only=True):
                rec = dict(zip(header, vals))
                if not rec.get("Azonosító"):
                    continue
                out = {"rater": rater, "item_id": rec["Azonosító"], "sheet": sheet}
                for hu, crit in rev.items():
                    v = rec.get(hu)
                    try:
                        out[f"h_{crit}"] = int(v) if v not in (None, "") else None
                    except (TypeError, ValueError):
                        out[f"h_{crit}"] = None
                wu = str(rec.get("Használnám (igen/nem)") or "").strip().lower()
                out["h_would_use"] = True if wu == "igen" else False if wu == "nem" else None
                out["comment"] = rec.get("Megjegyzés")
                rows.append(out)
    df = pd.DataFrame(rows).merge(key, on="item_id", how="left")
    missing = df["uid"].isna().sum()
    if missing:
        print(f"[rate-import] WARNING: {missing} rows have unknown item ids")
    blank = df[[f"h_{c}" for c in ("correctness", "clarity", "bloom_fit")]].isna().all(axis=1).sum()
    out = resolve(out_csv)
    out.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(out, index=False)
    print(f"[rate-import] {len(df)} ratings from {len(sheets)} sheets ({blank} left blank) -> {out}")
    return out
