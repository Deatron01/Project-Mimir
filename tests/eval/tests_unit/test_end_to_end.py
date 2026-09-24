"""Full harness run with fake services and a mock LLM: run -> score -> rate -> analyze."""
import json
import math
import random
import re

import pytest
from openpyxl import load_workbook

from mimir_eval import llm as llm_mod
from mimir_eval import runner as runner_mod
from mimir_eval.analysis import analyze
from mimir_eval.embed import BowEmbedder
from mimir_eval.rating import export_sheets, import_sheets
from mimir_eval.scoring import score_run
from mimir_eval.util import norm_text, read_jsonl


class FakeServices:
    """In-memory stand-in for Wellspring / RuneCarver / Bifrost."""
    fail_generate_for: set = set()

    def __init__(self, **_):
        self.store = []
        self.emb = BowEmbedder()

    def close(self):
        pass

    def health(self):
        return {"wellspring": "ok", "runecarver": "ok", "bifrost": "ok"}

    def extract(self, filename, data):
        return data.decode("utf-8"), 0.01

    def chunk(self, filename, ext, text, method="percentile", threshold=None, target_size=1000, encoder="window"):
        paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        return [{"type": "narrative", "content": p, "metadata": {}} for p in paras], 0.2

    def ingest(self, chunks):
        self.store = [c["content"] for c in chunks]
        return {"indexed_chunks": len(chunks)}, 0.05

    def search(self, query, limit):
        if not self.store:
            return [], 0.0
        q = self.emb.encode([query])
        m = self.emb.encode(self.store)
        order = (m @ q.T).ravel().argsort()[::-1][:limit]
        return [{"score": 0.5, "payload": {"text": self.store[i]}} for i in order], 0.02

    def generate(self, query, limit, fmt="json"):
        raise AssertionError("not used by naive_direct")


def mock_llm(messages, meta):
    system = messages[0]["content"]
    user = messages[-1]["content"]
    rng = random.Random(f"{meta['seed']}{len(user)}")
    if "exam reviewer" in system:                         # judge
        if "Answer this exam question" in user:
            letters = re.findall(r"^([A-D])\. ", user, re.M)
            return json.dumps({"answer": rng.choice(letters[:2]), "confidence": 0.8})
        s = rng.randint(2, 5)
        return json.dumps({"key_supported": s >= 3, "key_evidence": "x",
                           "distractors": [{"option": "B", "is_wrong": True, "on_topic": s > 2},
                                           {"option": "C", "is_wrong": True, "on_topic": True}],
                           "scores": {"correctness": s, "clarity": rng.randint(3, 5),
                                      "distractor_quality": rng.randint(2, 5), "bloom_fit": rng.randint(2, 5)},
                           "bloom_level": "understand", "would_use": s >= 4, "rationale": "ok"})
    # generator: build questions from context sentences
    n = int(re.search(r"pontosan (\d+) darab", user).group(1))
    ctx = user.split("KONTEXTUS:")[1].split("FELADAT")[0]
    sents = [s.strip() for s in re.split(r"(?<=[.!?])\s+", ctx) if len(s.strip()) > 20]
    k = n if meta["model"] == "good" else max(1, n - 2)   # the "bad" model returns too few questions
    qs = []
    for i in range(k):
        s = sents[i % len(sents)]
        qs.append({"type": "mcq", "text": f"Melyik állítás igaz: {s[:60]}?",
                   "answers": [{"text": s[:80], "is_correct": True}] +
                              [{"text": f"Hamis állítás {j} {rng.random():.3f}", "is_correct": False} for j in range(3)]})
    return json.dumps({"title": "Mimir AI Vizsga", "format": "pdf", "questions": qs}, ensure_ascii=False)


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.setattr(runner_mod, "MimirServices", FakeServices)
    llm_mod.set_mock_handler(mock_llm)
    cfgs = {}
    for name, model in (("A", "bad"), ("B", "good")):
        p = tmp_path / f"{name}.yaml"
        p.write_text(f"""name: {name}
documents: [hu-coffee, hu-immune, hu-gametheory]
seeds: [1, 2]
exam: {{n_questions: 4, types: [mcq]}}
pipeline: {{kind: naive_direct, retrieval_k: 2, probe_k: 3}}
generator: {{provider: mock, model: {model}}}
monitor_vram: false
""", encoding="utf-8")
        cfgs[name] = p
    yield tmp_path, cfgs
    llm_mod.set_mock_handler(None)


def test_full_pipeline(workspace):
    tmp, cfgs = workspace
    runs = {}
    for name, p in cfgs.items():
        rd = runner_mod.run_experiment(str(p), results_root=tmp / "results")
        exams = read_jsonl(rd / "exams.jsonl")
        assert len(exams) == 3 * 2 and all(e["status"] == "ok" for e in exams)
        assert all(e["retrieved"] and e["probe"] for e in exams)
        score_run(rd, judge_cfg={"provider": "mock", "model": "judge"}, embedder="bow")
        runs[name] = rd

    qa = read_jsonl(runs["A"] / "questions.jsonl")
    assert qa and {"grounded", "blind_correct", "judge_correctness", "meta_reference"} <= set(qa[0])
    # judge cache: re-scoring makes no new calls
    score_run(runs["A"], judge_cfg={"provider": "mock", "model": "judge"}, embedder="bow")
    assert json.loads((runs["A"] / "score.json").read_text())["judge_calls"] == 0

    # teacher rating round trip
    rdir = tmp / "ratings"
    export_sheets([str(runs["A"]), str(runs["B"])], str(rdir), n_per_run=6, n_raters=2, n_calibration=2)
    rng = random.Random(3)
    for sheet in sorted(rdir.glob("rater_*.xlsx")):
        wb = load_workbook(sheet)
        for ws_name in ("Kalibráció", "Értékelés"):
            ws = wb[ws_name]
            for row in ws.iter_rows(min_row=2):
                for c in row[10:14]:
                    c.value = rng.randint(2, 5)
                row[14].value = rng.choice(["igen", "nem"])
        wb.save(sheet)
    ratings = import_sheets([str(s) for s in sorted(rdir.glob("rater_*.xlsx"))],
                            str(rdir / "rating_key.csv"), str(rdir / "ratings.csv"))

    out = analyze([f"A={runs['A']}", f"B={runs['B']}"], "test", [("A", "B")], str(ratings),
                  out_root=str(tmp / "analysis"))
    summary = (out / "summary.md").read_text(encoding="utf-8")
    assert "Format compliance" in summary and "A vs B" in summary and "Krippendorff" in summary
    # the "bad" model never meets the requested count, the "good" one always does
    import pandas as pd
    q = pd.read_csv(out / "quality_by_arm.csv").set_index("column")
    assert q.loc["format_compliant", "A"] == 0.0 and q.loc["format_compliant", "B"] == 1.0
    for f in ("fig_quality_by_arm.png", "fig_latency_by_arm.png", "fig_judge_vs_teacher.png",
              "compare_A_vs_B.csv", "human.json"):
        assert (out / f).exists(), f
    # sheets must not leak which arm a question came from
    wb = load_workbook(next(rdir.glob("rater_*.xlsx")))
    text = " ".join(str(c.value) for row in wb["Értékelés"].iter_rows() for c in row if c.value)
    assert "|A|" not in text and runs["A"].name not in text
    assert norm_text("Értékelés")  # sanity


def test_resume_retries_failed(workspace, monkeypatch):
    tmp, cfgs = workspace
    calls = {"n": 0}

    def flaky(messages, meta):
        calls["n"] += 1
        if calls["n"] == 1:
            return "not json at all"
        return mock_llm(messages, meta)

    llm_mod.set_mock_handler(flaky)
    rd = runner_mod.run_experiment(str(cfgs["B"]), results_root=tmp / "results")
    statuses = [e["status"] for e in read_jsonl(rd / "exams.jsonl")]
    assert statuses.count("parse_error") == 1
    runner_mod.run_experiment(str(cfgs["B"]), resume=str(rd))
    from mimir_eval.scoring import latest_exams
    latest = latest_exams(read_jsonl(rd / "exams.jsonl"))
    assert len(latest) == 6 and all(e["status"] == "ok" for e in latest)


def test_report_sections(tmp_path, monkeypatch):
    """report: every TDK table (agreement, significance, languages, dataset) and the example questions."""
    from mimir_eval.report import build_report
    monkeypatch.setattr(runner_mod, "MimirServices", FakeServices)
    llm_mod.set_mock_handler(mock_llm)
    docs = "[hu-coffee, hu-immune, hu-gametheory, en-economics-short-01, en-law-short-07, en-lifesciences-short-04]"
    runs = []
    try:
        for name, model in (("E0", "bad"), ("E1", "good")):
            p = tmp_path / f"{name}.yaml"
            p.write_text(f"""name: {name}
documents: {docs}
seeds: [1]
exam: {{n_questions: 3, types: [mcq]}}
pipeline: {{kind: naive_direct, retrieval_k: 2, probe_k: 3}}
generator: {{provider: mock, model: {model}}}
monitor_vram: false
""", encoding="utf-8")
            rd = runner_mod.run_experiment(str(p), results_root=tmp_path / "results")
            score_run(rd, judge_cfg={"provider": "mock", "model": "judge"}, embedder="bow")
            runs.append(rd)
    finally:
        llm_mod.set_mock_handler(None)

    rdir = tmp_path / "ratings"
    export_sheets([str(r) for r in runs], str(rdir), n_per_run=6, n_raters=2, n_calibration=0)
    rng = random.Random(5)
    for sheet in sorted(rdir.glob("rater_*.xlsx")):
        wb = load_workbook(sheet)
        for row in wb["Értékelés"].iter_rows(min_row=2):
            for c in row[10:14]:
                c.value = rng.randint(1, 5)
            row[14].value = rng.choice(["igen", "nem"])
        wb.save(sheet)
    ratings = import_sheets([str(s) for s in sorted(rdir.glob("rater_*.xlsx"))],
                            str(rdir / "rating_key.csv"), str(rdir / "ratings.csv"))

    out = build_report(str(tmp_path / "results"), "t", out_root=str(tmp_path / "report"), ratings=str(ratings))
    for f in ("tabla3_egyetertes.png", "tabla4_szignifikancia.png", "tabla5_nyelvek.png", "abra7_nyelvek.png",
              "tabla6_adatkeszlet.png", "adatkeszlet.csv", "szignifikancia.csv", "peldak.md", "peldak.csv"):
        assert (out / f).exists(), f
    text = (out / "eredmenyek.md").read_text(encoding="utf-8")
    for h in ("A bíráló megbízhatósága", "Statisztikai próbák", "Magyar és angol", "Adatkészlet", "E1 vs E0"):
        assert h in text, h
    assert "PILOT" not in text                        # 6 documents: enough for the paired test
    import pandas as pd
    ds = pd.read_csv(out / "adatkeszlet.csv")
    assert ds["kiertekelve"].sum() == 6 and len(ds) > 6
    ex = pd.read_csv(out / "peldak.csv")
    assert set(ex["modszer"]) == {"E0", "E1"}
    # the paper's tables: one row per arm in a fixed order, CIs once there are enough documents
    main = (out / "latex" / "tab_main.tex").read_text(encoding="utf-8")
    assert "95\\,\\% CI" in main and main.index("\nE0 &") < main.index("\nE1 &")
    assert (out / "latex" / "tab_leakage.tex").exists() and not (out / "latex" / "tab_verifier.tex").exists()
    nums = json.loads((out / "latex" / "paper_numbers.json").read_text(encoding="utf-8"))
    assert [r["arm"] for r in nums["main"]] == ["E0", "E1"] and nums["n_documents"] == 6
    assert len((out / "latex" / "fig_quality_time.dat").read_text().splitlines()) == 3
    assert "E1 vs E0" in (out / "latex" / "tab_significance.tex").read_text(encoding="utf-8")

    # without ratings the agreement table is skipped, not an error
    out2 = build_report(str(tmp_path / "results"), "t2", out_root=str(tmp_path / "report"), ratings=None)
    assert not (out2 / "tabla3_egyetertes.png").exists()
    assert "még nem érkezett be" in (out2 / "eredmenyek.md").read_text(encoding="utf-8")


def test_pick_examples_prefers_grounded_and_blind_ok():
    import pandas as pd
    from mimir_eval.report import pick_examples
    base = {"arm": "E0", "judge_clarity": 4, "judge_distractor_quality": 4, "judge_bloom_fit": 4}
    qs = pd.DataFrame([
        {**base, "uid": "a", "grounded": True, "blind_correct": False, "judge_correctness": 5},
        {**base, "uid": "b", "grounded": True, "blind_correct": True, "judge_correctness": 4},
        {**base, "uid": "c", "grounded": False, "blind_correct": True, "judge_correctness": 5},
        {**base, "uid": "d", "grounded": True, "blind_correct": True, "judge_correctness": 2},
    ])
    p = pick_examples(qs)["E0"]
    assert p["jo"]["uid"] == "b"       # a has a higher score but failed the blind test
    assert p["rossz"]["uid"] == "a"    # a failed check (a, c) beats a merely low score (d); then lowest score


def test_judge_question_context_and_reasoning(tmp_path):
    from mimir_eval.judge import Judge
    seen = []
    llm_mod.set_mock_handler(lambda m, meta: seen.append(m) or json.dumps({"answer": "A"}))
    try:
        exam = {"doc_text": "D" * 100, "chunks": ["c0", "c1", "c2"], "context_text": "ctx"}
        j = Judge({"provider": "mock", "model": "j", "max_context_chars": 1000, "reasoning": "low",
                   "context": "question"}, tmp_path / "cache.jsonl")
        assert j.source_for(exam, {"context_ids": [2, 0]}) == ("c2\n\nc0", "question_context")
        assert j.source_for(exam, {}) == ("ctx", "retrieved_context")
        j._ask("blind", "p")
        assert seen[0][0]["content"].startswith("Reasoning: low\n")
        jd = Judge({"provider": "mock", "model": "j", "max_context_chars": 1000, "reasoning": None,
                    "context": "document"}, tmp_path / "cache2.jsonl")
        assert jd.source_for(exam, {"context_ids": [2]}) == ("D" * 100, "document")
        jd._ask("blind", "p")
        assert not seen[1][0]["content"].startswith("Reasoning")
    finally:
        llm_mod.set_mock_handler(None)


def test_full_document_baseline(tmp_path, monkeypatch):
    """B-doc: the whole (capped) document in one prompt; the cut share is recorded and scored."""
    monkeypatch.setattr(runner_mod, "MimirServices", FakeServices)
    llm_mod.set_mock_handler(mock_llm)
    try:
        p = tmp_path / "B-doc.yaml"
        p.write_text("""name: B-doc
documents: [hu-coffee]
seeds: [1]
exam: {n_questions: 3, types: [mcq]}
pipeline: {kind: full_document, max_doc_chars: 400, probe_k: 3}
generator: {provider: mock, model: good}
monitor_vram: false
""", encoding="utf-8")
        rd = runner_mod.run_experiment(str(p), results_root=tmp_path / "results")
        ex = read_jsonl(rd / "exams.jsonl")[0]
        assert ex["status"] == "ok" and len(ex["questions"]) == 3
        assert ex["retrieved"] == [] and ex["probe"]
        doc = ex["document"]
        assert doc["chars_used"] <= 400 < doc["chars"] and 0 < doc["truncated_share"] < 1
        assert ex["context_text"] and len(ex["context_text"]) == doc["chars_used"]
        score_run(rd, judge_cfg={"provider": "mock", "model": "judge"}, embedder="bow")
        import pandas as pd
        em = pd.read_csv(rd / "exam_metrics.csv")
        assert em["truncated_share"].iloc[0] == pytest.approx(doc["truncated_share"])
    finally:
        llm_mod.set_mock_handler(None)


def test_verifier_as_classifier():
    import pandas as pd
    from mimir_eval.paper_tables import verifier_rows
    rows = [("E2", True, True, True), ("E2", True, True, False), ("E2", True, False, True),
            ("E2", True, False, False), ("E2", True, True, True), ("E1", False, True, False)]
    qs = pd.DataFrame([{"arm": a, "verifier_on": on, "verified": v, "grounded": g, "blind_correct": None,
                        "judge_would_use": g} for a, on, v, g in rows])
    pooled, per_arm = verifier_rows(qs)
    assert [p["n"] for p in pooled] == [3, 2]                   # E1 ran without the verifier: left out
    assert pooled[0]["grounded"] == pytest.approx(2 / 3) and math.isnan(pooled[0]["blind_correct"])
    (r,) = per_arm
    assert r["arm"] == "E2" and r["precision"] == pytest.approx(2 / 3) and r["recall"] == pytest.approx(2 / 3)
    assert r["pass_rate"] == pytest.approx(0.6)
